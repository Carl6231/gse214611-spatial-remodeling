from pathlib import Path
from docx import Document
from docx.shared import Mm, Pt
import os,re,struct
root=Path(__file__).resolve().parents[2];out=Path(__file__).resolve().parents[1]
old=root/'10_全套图件精修_20261015/JTM_Manuscript_FiguresRefined_20261015.docx'
d=Document(old);original=Document(old)
captions={
1: 'Data structure and tissue identity. a, Study groups and analytical roles: eight formal-comparison tissues, two descriptive early timepoints and two auxiliary needle-pass tissues. b, H&E views of all 12 tissues, with retained-spot counts. c, Two mixed captures before and after author-provided barcode assignment; morphology was not used to infer identity. d, The initial six-class reference, including 28 Schwann-cell nuclei. e, Corrected primary-analysis accounting: 25,548 retained spots, 5,151 selected myeloid-rich centers, 3,964 assigned and 1,187 uncertain. The reference and spatial inputs share 22,490 genes. Complete H&E/spot pairs, QC distributions and marker compatibility are retained in Supplementary Figs. S1–S2. The original-reference 5,937-center analysis is shown separately in Supplementary Fig. S14. Spots are mixed transcriptomic measurements; physical tissues are the units of inference.',
2: 'Initial six-class deconvolution and method sensitivity. a, RCTD macrophage-labelled fractions in two Sham and six infarction tissues. All maps share a 0–1 scale with a square-root colour transform. b, Tissue-level macrophage-labelled load, rich-spot fraction and three major component fractions; dots denote physical tissues and bars denote group medians. The rich-area cutoff is 0.20. c, Composition context, displayed as within-class row z scores capped at ±2. d, RCTD–SpatialDWLS agreement and marker concordance, with stored Spearman coefficients shown numerically. e, Uncertain fractions across all 12 tissues. Sh1/2 denote Sham; D3/D7 denote infarction; 1 h/4 h are descriptive and N3/N4 auxiliary. These initial labels do not establish corrected macrophage purity, and the algorithms are not independent experiments. Complete profiles are retained in Supplementary Fig. S3.',
3: 'Reference correction and its spatial consequences. a, Reassignment of 18,438 originally macrophage-labelled nuclei; 8,693 (47.1%) remain macrophages. b, Original labels and corrected identities on identical saved UMAP coordinates in the 18,849-nucleus balanced reference. Other original lineages are grey; this balanced display is distinct from the full audited compartment counted in a. c, Marker evidence across the 12 audited clusters. Point area denotes the fraction of nuclei with detected transcript; colour denotes the saved log2 CPM ratio of a cluster to the remaining audited nuclei, with display limits capped at ±4. Cardiac transcripts in C0 remain compatible with ambient RNA alongside stromal markers and weak macrophage-core expression. d, Original and corrected macrophage fractions in three fixed representative tissues, using identical coordinates and a common 0–1 square-root colour scale. e, Paired tissue means for all 12 tissues; hollow and filled symbols denote original and corrected estimates. Reference balance, temporal support, calibration, complete broad maps and fine-state uncertainty are retained in Supplementary Figs. S6–S7.',
4: 'Rule-defined spatial neighborhoods under the corrected reference. a, Corrected niche assignments in all six replicated infarct tissues, retaining uncertain and nonselected spots. b, Within-tissue niche fractions among selected centers, with each tissue’s center count printed beside its bar. c, Four axes summarized by the median of tissue-level niche medians. d, One remodeling-axis summary per tissue. e, Same-label coherence above the stored permutation background; the 1-hour tissue has no center-to-center edges and is not evaluable (NE). f, Full contrasts and leave-one-tissue-out ranges, which are sensitivity ranges rather than confidence intervals. Selection used total myeloid fraction ≥0.20, radius multiplier 1.2, program gate ≥0.10 and margin gate ≥0.15. The analysis contains 5,151 centers, 3,964 assignments and 1,187 uncertain centers. Sh1/2 denote Sham, 1 h/4 h descriptive infarction and N3/N4 auxiliary needle-pass tissues.',
5: 'Day 7 expression structure and composition dependence. a, Saved joint-model R-squared values (purple) and residual variance (pale coral) for the fixed signature; labels also report five-spatial-block cross-validated R-squared. Residual variance is 2.3–4.1%, below the prespecified 25% criterion in all three tissues. b, Median center-spot signature scores, with one point per physical tissue. c, Three Day 7 spatial maps using the same symmetric display limits derived from the pooled 1st and 99th percentiles. d, Descriptive correlations for eight fixed genes excluded from axis construction. e, Nine stored pathway results; labels report normalized enrichment scores. Models A–D, reference-class summaries and axis-extreme decomposition are in Supplementary Fig. S9. The signature was selected against a composition-containing axis and does not establish an independent tissue state. Spatial CV is within tissue, not independent-animal validation.',
6: 'Gene recurrence, perturbation predictions and candidate-axis evidence. a, Sixteen nontechnical genes in the top 100 of all three Day 7 tissues. Colour denotes the stored spatial autocorrelation statistic; printed numbers denote within-tissue rank. Retained p-value and q-value fields were unavailable. b, Sdf4 and Abca1 maps selected by the fixed rank rule, with shared limits within each gene. c, Across-sample median CellOracle displacement for six targets, in model embedding units; all transition-mass summaries equal 0.500, and no target met the fixed module-consistency rule. d, Ten ordered ligand–receptor candidates with stage and tissue recurrence, minimum component coverage, neighborhood support, gradient support and raw CellChat scores. Filled/open symbols mean supported/unsupported; NE means not evaluable. Candidate axes retain their original communication labels, while simulated perturbations remain model-conditional predictions. Complete evidence is retained in Supplementary Figs. S10–S13.',
7: 'Context-specific mixed-spot associations and evaluability. a, Dominant RCTD8 classes in six infarct tissues. b, Both directions of the Day 7 mononuclear-myeloid–fibroblast relationship in transition and injury-core contexts; each line connects the two context estimates from one tissue and does not indicate a temporal trajectory. c, Complete contextual L estimates at the fixed 200-µm radius. A cross marks an absent class and a dash an absent region; missing values are not zero effects. d, Evaluable and unavailable records across 22 tissues, ordered as formal comparisons, descriptive/auxiliary observations and the external cohort. External BANKSY domains have within-tissue meanings only. CM denotes cardiomyocyte, Fib fibroblast, Myel mononuclear myeloid and Neu neutrophil. These mixed-spot associations do not establish cell contact, mechanism or statistical significance. Complete tissue-specific results appear in Supplementary Fig. S15.',
8: 'Secreted-protein activity and signaling patterns under the fixed SecAct model. a, Candidate-ligand activity in the myeloid-labelled single-nucleus subset and spatial tissue medians, using separate symmetric activity scales. SRR suffixes identify nucleus samples; the modalities are not animal-paired. b, Fractions of modeled spots with sending and receiving vectors, shown for all 72 tissue–ligand–direction combinations on a shared 0–1 scale. c, Zero retained candidate-ligand edges in every one of the 36 ligand–nucleus-sample combinations, considering all source–receiver class pairs. d, Numbers of independently selected spatial factors and retained proteins per tissue. Activity, edges and vector coverage are model outputs, not measured secretion, receptor activation or movement. Factor identities are local to each tissue. Complete activity, communication, pattern and paired vector displays are retained in Supplementary Figs. S16–S18.'}
replacements={
75:('Fig. 2a–e','Fig. 2a–d'),76:('Fig. 2f–h','Fig. 2d,e'),78:('Fig. 3a,b','Fig. 3a–c'),
79:('(Fig. 3c–f).','(Supplementary Fig. S6; Fig. 3d,e).'),
80:('(Fig. 3g,h and Supplementary Figs. S6 and S7)','(Supplementary Figs. S6 and S7)'),
82:('Fig. 4a–d','Fig. 4a,b'),83:('Fig. 4e,f','Fig. 4b,e'),84:('Fig. 4g,h','Fig. 4d,f'),
86:('Fig. 5a–d','Fig. 5b–e'),87:('Fig. 5e','Fig. 5a'),95:('Fig. 6c','Fig. 6d'),101:('Fig. 6d','Fig. 6c'),
107:('shows tissue-level means and the largest reference component per spot','shows tissue means, maximum-coefficient distributions and spot-wise largest components')}
def replace_para(p,text):
 if p.runs:
  p.runs[0].text=text
  for r in p.runs[1:]:r.text=''
 else:p.add_run(text)
for i,(a,b) in replacements.items():
 p=d.paragraphs[i];assert a in p.text,(i,a);replace_para(p,p.text.replace(a,b))
for i,p in enumerate(d.paragraphs):
 m=re.match(r'^Figure ([1-8])\.',p.text)
 if m:
  n=int(m[1]);replace_para(p,f'Figure {n}. {captions[n]}')
  for r in p.runs:r.font.size=Pt(10)
  p.paragraph_format.line_spacing=1.05
  p.paragraph_format.keep_together=True
  p.paragraph_format.space_before=Pt(6)
  p.paragraph_format.space_after=Pt(0)
for i,shape in enumerate(d.inline_shapes,1):
 data=(out/'Figures'/f'Figure_{i}.png').read_bytes();assert data[:8]==b'\x89PNG\r\n\x1a\n'
 w,h=struct.unpack('>II',data[16:24]);rid=shape._inline.graphic.graphicData.pic.blipFill.blip.embed
 d.part.related_parts[rid]._blob=data;shape.width=Mm(160);shape.height=Mm(160*h/w)
 shape._inline.docPr.set('descr',f'Figure {i}. {captions[i].split(". ")[0]}.')
commit=os.environ.get('JTM_CODE_COMMIT')
if commit:
 assert re.fullmatch('[0-9a-f]{40}',commit)
 for p in d.paragraphs:
  for r in p.runs:r.text=re.sub(r'/tree/[0-9a-f]{40}', '/tree/'+commit,r.text)
assert len(d.inline_shapes)==8 and len(d.tables)==3
assert [p.text for p in original.paragraphs[:18]]==[p.text for p in d.paragraphs[:18]]
assert [[c.text for row in t.rows for c in row.cells] for t in original.tables]==[[c.text for row in t.rows for c in row.cells] for t in d.tables]
assert len([p for p in d.paragraphs if re.match(r'^\d+\. ',p.text)])==63
assert 'w:lnNumType' in d.sections[0]._sectPr.xml
path=out/'JTM_Manuscript_Redesigned_20261015.docx';d.save(path)
print('Saved',path.name,'8 proportion-preserved figures; 63 references; author block and all three tables unchanged')
